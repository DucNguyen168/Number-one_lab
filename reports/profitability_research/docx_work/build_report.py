from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[3]
REFERENCE = Path(r"C:\Users\Admin\.codex\plugins\cache\openai-curated-remote\openai-templates\0.1.1\skills\artifact-template-legal-memorandum\assets\reference.docx")
OUTPUT = ROOT / "reports" / "profitability_research" / "KE_HOACH_NGHIEN_CUU_LOI_NHUAN_60_80_SO.docx"


def font(run, name="Garamond", size=10.5, bold=None, color="000000", italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return run


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_table_header(row):
    repeat_header(row)
    no_split(row)


def set_table_borders(tbl):
    tbl_pr = tbl._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "B8C1CC")
        borders.append(node)


def set_paragraph_keep(paragraph, keep_next=False, keep_lines=True):
    p_pr = paragraph._p.get_or_add_pPr()
    if keep_next:
        p_pr.append(OxmlElement("w:keepNext"))
    if keep_lines:
        p_pr.append(OxmlElement("w:keepLines"))


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "365F91")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_fonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "eastAsia"):
        r_fonts.set(qn(f"w:{attr}"), "Garamond")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "21")
    r_pr.extend([r_fonts, color, underline, size])
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def clear_body(document):
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def p(document, text="", *, bold_prefix=None, indent=True, after=5, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    para = document.add_paragraph()
    para.alignment = align
    para.paragraph_format.space_after = Pt(after)
    para.paragraph_format.line_spacing = 1.08
    if indent:
        para.paragraph_format.first_line_indent = Inches(0.3)
    if bold_prefix and text.startswith(bold_prefix):
        font(para.add_run(bold_prefix), bold=True)
        font(para.add_run(text[len(bold_prefix):]), italic=italic)
    else:
        font(para.add_run(text), italic=italic)
    return para


def bullet(document, text, number=None, level=0):
    para = document.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.3 + level * 0.25)
    para.paragraph_format.first_line_indent = Inches(-0.22)
    para.paragraph_format.space_after = Pt(3)
    marker = f"{number}." if number is not None else "•"
    font(para.add_run(marker + " "), bold=number is not None)
    font(para.add_run(text))
    set_paragraph_keep(para)
    return para


def heading(document, text, level=1):
    para = document.add_paragraph(style=f"Heading {level}")
    para.add_run(text)
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    para.paragraph_format.space_after = Pt(5)
    for run in para.runs:
        font(run, size=13 if level == 1 else 11.5, bold=True, color="000000")
        run.underline = level == 1
    set_paragraph_keep(para, keep_next=True)
    return para


def table(document, headers, rows, widths=None):
    tbl = document.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    set_table_borders(tbl)
    hdr = tbl.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        shade(cell, "263B52")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font(para.add_run(str(text)), size=9, bold=True, color="FFFFFF")
        if widths:
            cell.width = Inches(widths[i])
    for ri, values in enumerate(rows):
        row = tbl.add_row()
        no_split(row)
        for i, value in enumerate(values):
            cell = row.cells[i]
            set_cell_margins(cell)
            if ri % 2:
                shade(cell, "F2F4F7")
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            font(para.add_run(str(value)), size=8.8)
            if widths:
                cell.width = Inches(widths[i])
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return tbl


def callout(document, title, text):
    para = document.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.12)
    para.paragraph_format.right_indent = Inches(0.12)
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after = Pt(8)
    p_pr = para._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "EAF0F6")
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), "365F91")
    borders.append(left)
    p_pr.append(borders)
    font(para.add_run(title + "  "), bold=True, color="263B52")
    font(para.add_run(text))
    set_paragraph_keep(para)


def source(document, number, title, citation, url):
    para = document.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.28)
    para.paragraph_format.first_line_indent = Inches(-0.28)
    para.paragraph_format.space_after = Pt(4)
    font(para.add_run(f"[{number}] "), bold=True)
    add_hyperlink(para, title, url)
    font(para.add_run(". " + citation), size=9.5)


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(REFERENCE, OUTPUT)
    doc = Document(OUTPUT)
    clear_body(doc)

    styles = doc.styles
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = styles[style_name]
        style.font.name = "Garamond" if style_name != "Title" else "Play"
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), style.font.name)
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), style.font.name)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), style.font.name)
        style.font.color.rgb = RGBColor(0, 0, 0)
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.size = Pt(21)
    styles["Title"].font.bold = True

    header = doc.sections[0].header
    for para in header.paragraphs:
        for run in para.runs:
            run.text = ""
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(hp.add_run("BÁO CÁO NGHIÊN CỨU · 05/09/2026"), name="Arial", size=7.5, color="A6A6A6")

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(22)
    title.paragraph_format.space_after = Pt(18)
    font(title.add_run("KẾ HOẠCH NGHIÊN CỨU\nTỐI ƯU LỢI NHUẬN DÀN 60 VÀ 80 SỐ"), name="Play", size=21, bold=True)

    metadata = [
        ("Mục tiêu", "Tối đa hóa lợi nhuận kỳ vọng có kiểm định, không tối ưu riêng kết quả lịch sử"),
        ("Dữ liệu", "XSMB hai số cuối từ 01/01/2019 đến 05/09/2026"),
        ("Mức mô phỏng", "10.000 đồng mỗi số; trúng nhận tổng cộng 999.000 đồng"),
        ("Trạng thái", "Kế hoạch nghiên cứu; chưa chứng minh có lợi thế tương lai"),
    ]
    for label, value in metadata:
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Inches(0.15)
        para.paragraph_format.space_after = Pt(3)
        font(para.add_run(label + ":\t"), bold=True)
        font(para.add_run(value))
    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(8)
    rule.paragraph_format.space_after = Pt(10)
    p_pr = rule._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "263B52")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    heading(doc, "Kết luận điều hành")
    p(doc, "Trong lưới chiến lược đã chạy, ứng viên tốt nhất về lợi nhuận tuyệt đối ở 361 kỳ quan sát gần nhất là dàn 60, loại 40 số khác nhau vừa xuất hiện, tham gia mỗi kỳ: đúng 234/361 (64,82%), lãi mô phỏng 17.166.000 đồng, ROI 7,93%, drawdown tối đa 4.824.000 đồng.")
    p(doc, "Ứng viên tốt nhất về ROI là dàn 60, loại 38 số khác nhau vừa xuất hiện, nghỉ sau một kỳ sai và chỉ tham gia lại sau một kỳ theo dõi đúng: 154/232 lần tham gia đúng (66,38%), lãi 14.646.000 đồng, ROI 10,52%, drawdown 4.410.000 đồng. Dàn 80 tốt nhất là loại 20 số vừa xuất hiện và tham gia mỗi kỳ: 300/361 (83,10%), lãi 10.900.000 đồng, ROI 3,77%, drawdown 3.455.000 đồng.")
    callout(doc, "Quyết định hiện tại", "Giữ ba ứng viên 60/38/nghỉ, 60/40/mỗi kỳ và 80/20/mỗi kỳ để kiểm định. Chưa dùng kết quả thăm dò làm cam kết không lỗ. Ưu tiên xây registry, nested walk-forward và mô phỏng toàn pipeline trước khi thêm mô hình phức tạp.")
    p(doc, "Các kết quả này là thăm dò. Lưới loại trừ và dữ liệu đã được xem nhiều lần. White chỉ ra rằng dùng lại cùng dữ liệu cho lựa chọn mô hình và suy luận có thể tạo kết quả đẹp do may mắn; Reality Check kiểm định mô hình tốt nhất trong cả tập ứng viên so với chuẩn [1].")

    heading(doc, "1 Bài toán kinh tế phải tối ưu")
    p(doc, "Với dàn N số, lợi nhuận một kỳ tham gia là 999.000 × H − 10.000 × N, trong đó H bằng 1 nếu trúng và 0 nếu sai. Điểm hòa vốn là 60,0601% cho dàn 60 và 80,0801% cho dàn 80.")
    p(doc, "Nếu hai dàn dùng cùng thứ hạng và dàn 80 chứa dàn 60, chỉ nên mua thêm 20 số khi xác suất thật cộng thêm của chúng vượt 200.000 / 999.000 = 20,0200%.")
    p(doc, "Mục tiêu chọn mô hình là lợi nhuận trung bình trên mỗi kỳ quan sát, trong đó ngày nghỉ ghi 0 đồng. ROI, hit rate, số kỳ tham gia, drawdown tối đa và chuỗi thua là chỉ số phụ. Cách này ngăn một quy tắc nghỉ rất nhiều ngày trông tốt chỉ vì mẫu số nhỏ.")

    heading(doc, "2 Dữ liệu và phép thử đã thực hiện")
    for text in [
        "Snapshot có 2.752 kết quả hợp lệ từ 01/01/2019 đến 05/09/2026; 53 ngày lịch không có kết quả trong nguồn và không được tự điền.",
        "SHA-256 của snapshot là 8078474323e36520602e02a86692db7d73ff7bbd434f28e09ec8ca084c13ef37.",
        "Mỗi dự đoán chỉ dùng dữ liệu đứng trước ngày cần dự đoán. 365 kết quả đầu làm lịch sử tối thiểu; còn 2.387 kỳ walk-forward từ 05/01/2020.",
        "Lưới: dàn 60 với K = 0, 5, 10, 20, 30, 38, 40; dàn 80 với K = 0, 5, 10, 20; mỗi cấu hình chạy hằng ngày và chạy với quy tắc nghỉ sau sai.",
        "Loại K nghĩa là bỏ K số khác nhau xuất hiện gần nhất, rồi lấy các số có thứ hạng tiếp theo cho đủ N. K = 40, N = 60 lấy toàn bộ 60 số chưa nằm trong nhóm 40 số gần nhất.",
    ]:
        bullet(doc, text)

    heading(doc, "3 Kết quả thăm dò")
    heading(doc, "3 1 Cửa sổ 361 kỳ quan sát gần nhất", 2)
    table(doc,
          ["Chiến lược", "Đúng cược", "Tỷ lệ", "Lãi ròng", "ROI", "Drawdown"],
          [
              ["60 loại 40 mỗi kỳ", "234 / 361", "64,82%", "+17.166.000", "7,93%", "4.824.000"],
              ["60 loại 38 nghỉ", "154 / 232", "66,38%", "+14.646.000", "10,52%", "4.410.000"],
              ["60 loại 38 mỗi kỳ", "231 / 361", "63,99%", "+14.169.000", "6,54%", "6.228.000"],
              ["80 loại 20 mỗi kỳ", "300 / 361", "83,10%", "+10.900.000", "3,77%", "3.455.000"],
              ["80 loại 20 nghỉ", "251 / 301", "83,39%", "+9.949.000", "4,13%", "5.015.000"],
              ["Top 60 hiện tại", "191 / 361", "52,91%", "−25.791.000", "−11,91%", "28.173.000"],
              ["Top 80 hiện tại", "277 / 361", "76,73%", "−12.077.000", "−4,18%", "16.924.000"],
          ], [1.75, .75, .65, 1.0, .65, .9])
    p(doc, "Nếu giả sử đây là một chiến lược đã khóa trước và các lần trúng phù hợp giả định Bernoulli, dàn 60 loại 40 có p-value nhị thức một phía khoảng 0,0358. Vì phương án được chọn sau khi so sánh nhiều K, N và quy tắc tham gia, con số đó không phải p-value xác nhận.")

    heading(doc, "3 2 Toàn bộ 2 387 kỳ walk forward", 2)
    table(doc,
          ["Chiến lược", "Đúng cược", "Tỷ lệ", "Lãi ròng", "ROI", "Drawdown"],
          [
              ["60 loại 38 nghỉ", "883 / 1.446", "61,07%", "+14.517.000", "1,67%", "29.967.000"],
              ["60 loại 38 mỗi kỳ", "1.445 / 2.387", "60,54%", "+11.355.000", "0,79%", "35.037.000"],
              ["60 loại 40 mỗi kỳ", "1.443 / 2.387", "60,45%", "+9.357.000", "0,65%", "34.647.000"],
              ["80 loại 20 mỗi kỳ", "1.905 / 2.387", "79,81%", "−6.505.000", "−0,34%", "33.588.000"],
              ["Top 60 hiện tại", "1.384 / 2.387", "57,98%", "−49.584.000", "−3,46%", "64.674.000"],
              ["Top 80 hiện tại", "1.890 / 2.387", "79,18%", "−21.490.000", "−1,13%", "33.121.000"],
          ], [1.75, .75, .65, 1.0, .65, .9])
    p(doc, "Chiến lược 60/38/nghỉ lỗ trong 2020–2023 và có lãi trong 2024–2026. Khoảng Wilson 95% toàn kỳ cho hit rate là 58,53–63,55%, vẫn chứa mốc hòa vốn 60,0601%; kiểm định nhị thức một phía chưa hiệu chỉnh có p xấp xỉ 0,226. Mẫu gần đây có thể là thay đổi chế độ hoặc dao động ngẫu nhiên.")
    p(doc, "Một mô phỏng chọn chiến lược bằng các năm trước rồi áp dụng cho năm kế tiếp chọn không tham gia trong 2024 và 2025; đến 2026 mới chọn 60/38/nghỉ và tạo lãi mô phỏng 9.894.000 đồng trên 160 lần tham gia. Đây vẫn là tái dựng sau khi lịch sử đã được nghiên cứu, không phải holdout sạch.")

    heading(doc, "4 Cách hiểu phương pháp loại trừ")
    p(doc, "Nghiên cứu về hành vi người chơi xổ số cho thấy người chơi có xu hướng tránh số vừa ra; điều đó không biến lượt quay độc lập thành quy luật số vừa ra ít có khả năng ra lại [2]. Loại số gần đây phải được xem như một giả thuyết dữ liệu.")
    p(doc, "Với một phân phối xác suất đã ước lượng đúng, chọn N số có xác suất cao nhất tối đa hóa lợi nhuận kỳ vọng. Cách loại trừ tốt nhất khi đó là loại các số có xác suất thấp nhất. Kế hoạch mới biến độ gần đây thành đặc trưng hoặc mức phạt có regularization; K = 38 và 40 vẫn được giữ làm đối chứng cố định.")

    heading(doc, "5 Danh mục thuật toán nghiên cứu")
    heading(doc, "5 1 Đối chứng bắt buộc", 2)
    for i, text in enumerate([
        "Phân phối đều 1% cho mỗi số và không tham gia, lợi nhuận 0.",
        "frequency_ensemble_v1 hiện tại.",
        "60/38/nghỉ, 60/40/mỗi kỳ và 80/20/mỗi kỳ, đóng băng đúng định nghĩa đã chạy.",
    ], 1):
        bullet(doc, text, i)

    heading(doc, "5 2 Mô hình ưu tiên", 2)
    models = [
        ("Bayesian frequency có co về đều", "p(j) = (count(j) + A/100) / (n + A). Dirichlet là prior liên hợp của multinomial và posterior cập nhật bằng số quan sát [3]. Chỉ thử lưới nhỏ A và vài half-life đăng ký trước."),
        ("Mô hình tách hàng chục và hàng đơn vị", "Trộn xác suất 100 lớp trực tiếp với tích hai phân phối 10 lớp: p(a,b) = λ p_direct(a,b) + (1−λ) p_tens(a)p_units(b)."),
        ("Markov bậc một có co mạnh", "Ma trận 100×100 có 10.000 ô nhưng khoảng 2.751 chuyển tiếp. Mỗi hàng co về phân phối nền; chưa thử Markov bậc cao."),
        ("Gradient boosting nhỏ", "Dùng một thư viện, độ sâu thấp và ngân sách tham số cố định. Benchmark dữ liệu bảng cho thấy cây là đối chứng mạnh ở mẫu cỡ trung bình, nhưng không chứng minh hiệu quả xổ số [4]."),
        ("TabPFN-3 làm challenger duy nhất", "Báo cáo 2026 công bố kết quả mạnh trên benchmark nhiều lớp [5]; tài liệu nhà cung cấp nêu giới hạn 160 lớp, đủ về mặt kỹ thuật cho 100 nhãn [6]."),
    ]
    for i, (name, detail) in enumerate(models, 1):
        para = bullet(doc, "", i)
        para.runs[-1].text = ""
        font(para.add_run(name + ". "), bold=True)
        font(para.add_run(detail))
    p(doc, "Không ưu tiên LSTM hoặc Transformer tự huấn luyện: 2.752 nhãn quá ít so với độ linh hoạt và chi phí thử tham số. Chỉ đưa vào nếu kiểm định phụ thuộc thời gian cho thấy tín hiệu ổn định và có thêm dữ liệu.")

    heading(doc, "5 3 Đặc trưng đã khóa", 2)
    for text in [
        "Tần suất ở cửa sổ 30, 90, 180, 365 và toàn lịch sử; tần suất suy giảm mũ với tối đa ba half-life.",
        "Số kỳ kể từ lần xuất hiện; cờ xuất hiện trong 1, 3, 5, 10, 20, 38 và 40 số khác nhau gần nhất.",
        "Tần suất hàng chục, hàng đơn vị, cặp chữ số và chuyển tiếp từ kết quả trước có prior.",
        "Thứ trong tuần làm nhóm kiểm tra riêng; loại nếu không cải thiện outer walk-forward.",
    ]:
        bullet(doc, text)
    p(doc, "Mọi mô hình phải xuất đủ 100 xác suất có tổng bằng 1. Log score và Brier score là proper scoring rules để đánh giá dự báo xác suất [7]. Temperature scaling fit trên validation có thể cải thiện hiệu chỉnh [8], nhưng giữ nguyên thứ hạng nên không tự tăng hit rate top 60/80; giá trị chính là hỗ trợ quyết định nghỉ.")

    heading(doc, "6 Thiết kế kiểm định tránh tự lừa mình")
    steps = [
        "Sổ đăng ký thử nghiệm. Mỗi N, K, cửa sổ, half-life, feature set, thuật toán, seed, quy tắc nghỉ và ngưỡng tham gia là một trial. Không xóa trial thua.",
        "Nested walk-forward. Tầng trong chọn tham số bằng quá khứ; tầng ngoài tiếp theo ghi dự đoán và lợi nhuận. Chọn tham số và đánh giá trên cùng dữ liệu tạo ước lượng quá lạc quan [9].",
        "Chuẩn kinh tế. So với không tham gia = 0 đồng; so với phân phối đều để kiểm tra năng lực dự báo. Thắng mô hình hiện tại đang lỗ chưa đủ chứng minh có lãi.",
        "Sửa nhiều phép thử. Chạy Reality Check và SPA trên ma trận lợi nhuận cùng ngày của toàn bộ trial [1,10]. Bootstrap lấy cùng chỉ số ngày cho mọi chiến lược.",
        "Mô phỏng không tín hiệu. Tạo tối thiểu 10.000 chuỗi IID đều 00–99, giữ độ dài và ngày thiếu, rồi chạy lại toàn bộ xây đặc trưng, chọn mô hình và đánh giá.",
        "Theo dõi tương lai. Khóa một champion và hash prediction trước giờ quay. Confidence sequence phù hợp với theo dõi liên tục [11], nhưng không sửa thiên lệch lựa chọn từ lịch sử.",
    ]
    for i, text in enumerate(steps, 1):
        bullet(doc, text, i)

    heading(doc, "7 Quy tắc chọn 60 hay 80 và quyết định tham gia")
    p(doc, "Mỗi ngày, mô hình tạo p(00)…p(99), hiệu chỉnh bằng validation quá khứ và xếp giảm dần. q60 là tổng xác suất top 60; q80 là tổng xác suất top 80.")
    for text in [
        "Dàn 60 đủ điều kiện nếu cận dưới đã hiệu chỉnh của q60 vượt 0,600601 cộng biên an toàn đã khóa.",
        "Dàn 80 chỉ thay dàn 60 nếu cận dưới của q80 − q60 vượt 0,200200 và lợi nhuận kỳ vọng trên mỗi kỳ quan sát cao hơn sau chi phí.",
        "Nếu không điều kiện nào đạt, chỉ ghi prediction và không tham gia.",
        "Không dùng chuỗi thắng thua riêng lẻ để đổi tiền. Quy tắc nghỉ chỉ tồn tại nếu qua nested walk-forward và kiểm định nhiều mô hình.",
    ]:
        bullet(doc, text)
    p(doc, "Với một chiến lược cố định và giả định Bernoulli phù hợp, 365 lần tham gia cần ít nhất 236 lần đúng cho dàn 60 hoặc 306 lần đúng cho dàn 80 để vượt hòa vốn trong kiểm định nhị thức một phía 5%. Mốc này không áp dụng cho phương án được chọn sau nhiều phép thử.")
    table(doc, ["Lợi thế thật trên hòa vốn", "60 số", "80 số"], [
        ["+1 điểm phần trăm", "≈ 14.787 kỳ", "≈ 9.734 kỳ"],
        ["+2 điểm phần trăm", "≈ 3.685 kỳ", "≈ 2.400 kỳ"],
        ["+3 điểm phần trăm", "≈ 1.632 kỳ", "≈ 1.051 kỳ"],
    ], [2.6, 1.8, 1.8])
    p(doc, "Đây là xấp xỉ power 80% cho một kiểm định một phía 5%, chưa tăng thêm vì phụ thuộc, nhiều phép thử hoặc ngày nghỉ. Một năm có rất ít lực kiểm định cho lợi thế nhỏ.")

    heading(doc, "8 Kế hoạch triển khai cụ thể")
    phases = [
        ["0 Đóng băng", "1–2 ngày", "Snapshot, checksum, ngày thiếu, payout; tái tạo comparison.csv và results.json"],
        ["1 Pipeline và registry", "3–5 ngày", "Interface fit(history) → p[100]; trial registry; test chống leakage"],
        ["2 Thống kê gọn", "1 tuần", "Bayesian shrinkage, decay, digit mixture, Markov co mạnh; tối đa 12 cấu hình inner"],
        ["3 Challenger ML", "1 tuần", "Một boosting nhỏ; một TabPFN-3 nếu khả thi; calibration chỉ bằng quá khứ"],
        ["4 Kiểm định", "1–2 tuần", "Nested walk-forward, Reality Check/SPA, 10.000 mô phỏng toàn pipeline, model card"],
        ["5 Paper trading", "≥ 365 kỳ", "Khóa champion và cấu hình; prediction trước giờ quay; thay đổi tạo phiên bản mới"],
    ]
    table(doc, ["Giai đoạn", "Thời lượng", "Đầu ra bắt buộc"], phases, [1.35, .9, 4.05])
    p(doc, "Champion được chọn theo lợi nhuận trung bình trên mỗi kỳ quan sát. ROI, drawdown và độ ổn định là điều kiện phụ đã đăng ký trước; không tạo điểm tổng hợp tùy ý sau khi xem kết quả.")

    heading(doc, "9 Điều kiện vượt vòng và điều kiện dừng")
    gates = [
        "Lợi nhuận outer walk-forward dương so với không tham gia.",
        "Cận dưới một phía của lợi nhuận kỳ vọng vượt 0 sau điều chỉnh tìm kiếm theo kiểm định đã khóa.",
        "Reality Check, SPA hoặc mô phỏng toàn pipeline bác bỏ hợp lý việc phương án tốt nhất chỉ do thử nhiều cách.",
        "Paper trading khóa trước đạt tiêu chí mẫu và rủi ro đã định.",
        "Không có leakage, sửa dữ liệu hồi tố hoặc thay đổi payout.",
    ]
    for i, text in enumerate(gates, 1):
        bullet(doc, text, i)
    p(doc, "Dừng hoặc quay về chỉ theo dõi nếu cận dưới không vượt 0, pipeline không thắng phân phối tốt nhất trên dữ liệu IID giả, drawdown vượt giới hạn đăng ký, hoặc lợi thế mất ở hai outer window liên tiếp. Kelly tối ưu tăng trưởng vốn khi đã có thông tin dự đoán và odds; nó không tạo lợi thế cho mô hình kỳ vọng âm [12]. Trong nghiên cứu giữ mức mô phỏng 10.000 đồng mỗi số và không gấp thếp.")

    heading(doc, "10 Khuyến nghị hiện tại")
    p(doc, "Giữ ba ứng viên đóng băng: 60/38/nghỉ, 60/40/mỗi kỳ và 80/20/mỗi kỳ. Ưu tiên 60/38/nghỉ vì ROI toàn kỳ dương và vốn tham gia thấp hơn; dùng 60/40/mỗi kỳ làm challenger lợi nhuận tuyệt đối gần đây; dùng 80/20/mỗi kỳ làm đối chứng biến động thấp. Không chọn ứng viên chỉ dựa trên 361 kỳ gần nhất.")
    callout(doc, "Bước kế tiếp có giá trị cao nhất", "Xây experiment registry, nested walk-forward và mô phỏng toàn pipeline trước khi thêm mô hình phức tạp. Nếu chuỗi không có tín hiệu bền vững, thuật toán mạnh hơn sẽ chỉ học nhiễu nhanh hơn.")

    heading(doc, "Nguồn")
    sources = [
        ("A Reality Check for Data Snooping", "Halbert White, Econometrica 68(5), 2000", "https://users.ssc.wisc.edu/~behansen/718/White2000.pdf"),
        ("The Gambler’s Fallacy in Lottery Play", "Charles Clotfelter and Philip Cook, NBER Working Paper 3769, 1991; Management Science, 1993", "https://www.nber.org/papers/w3769"),
        ("Time-Series Models — Analytic posterior", "Stan User’s Guide 2.35", "https://mc-stan.org/docs/2_35/stan-users-guide/time-series.html#analytic-posterior"),
        ("Why do tree-based models still outperform deep learning on typical tabular data", "Léo Grinsztajn, Edouard Oyallon and Gaël Varoquaux, NeurIPS 2022", "https://arxiv.org/abs/2207.08815"),
        ("TabPFN-3 Technical Report", "Grinsztajn et al., arXiv:2605.13986, revised 28/05/2026", "https://arxiv.org/abs/2605.13986"),
        ("REST API quickstart — model limits", "Prior Labs, accessed 05/09/2026", "https://docs.priorlabs.ai/api-reference/getting-started"),
        ("Strictly Proper Scoring Rules Prediction and Estimation", "Tilmann Gneiting and Adrian Raftery, JASA 102(477), 2007", "https://sites.stat.washington.edu/people/raftery/Research/PDF/Gneiting2007jasa.pdf"),
        ("On Calibration of Modern Neural Networks", "Chuan Guo et al., ICML 2017", "https://proceedings.mlr.press/v70/guo17a.html"),
        ("Cross-validation evaluating estimator performance", "scikit-learn official documentation, accessed 05/09/2026", "https://scikit-learn.org/stable/modules/cross_validation.html"),
        ("A Test for Superior Predictive Ability", "Peter R. Hansen, Journal of Business & Economic Statistics 23(4), 2005", "https://bashtage.github.io/kevinsheppard.com/files/teaching/mfe/advanced-econometrics/Hansen.pdf"),
        ("Time-uniform nonparametric nonasymptotic confidence sequences", "Steven Howard et al., Annals of Statistics 49(2), 2021; revised 06/08/2022", "https://arxiv.org/abs/1810.08240"),
        ("A New Interpretation of Information Rate", "John L. Kelly Jr., Bell System Technical Journal 35(4), 1956", "https://onlinelibrary.wiley.com/doi/abs/10.1002/j.1538-7305.1956.tb03809.x"),
    ]
    for i, (title_text, citation, url) in enumerate(sources, 1):
        source(doc, i, title_text, citation, url)

    doc.core_properties.title = "Kế hoạch nghiên cứu tối ưu lợi nhuận dàn 60 và 80 số"
    doc.core_properties.subject = "Backtest, thuật toán, kiểm định và kế hoạch paper trading XSMB"
    doc.core_properties.author = "XSMB Research Lab"
    doc.core_properties.keywords = "XSMB, walk-forward, dàn 60, dàn 80, loại trừ, lợi nhuận"
    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
